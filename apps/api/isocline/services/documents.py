"""File validation, text extraction and chunking. Uploaded files are parsed, never executed."""
from __future__ import annotations

import csv
import io

ALLOWED = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    ".csv": {"text/csv", "application/vnd.ms-excel", "text/plain"},
    ".json": {"application/json", "text/plain"},
}


class DocumentError(ValueError):
    pass


def validate_upload(filename: str, content_type: str | None, data: bytes, max_mb: int) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise DocumentError(f"Unsupported file type '{ext or 'none'}'. Allowed: {', '.join(sorted(ALLOWED))}")
    if len(data) > max_mb * 1024 * 1024:
        raise DocumentError(f"File exceeds the {max_mb} MB limit")
    if not data:
        raise DocumentError("File is empty")
    # Magic-byte checks: content must match the claimed type.
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise DocumentError("File content is not a valid PDF")
    if ext == ".docx" and not data.startswith(b"PK"):
        raise DocumentError("File content is not a valid DOCX")
    if ext in (".txt", ".md", ".csv", ".json"):
        try:
            data[:65536].decode("utf-8")
        except UnicodeDecodeError as e:
            raise DocumentError("Text files must be UTF-8") from e
        if b"\x00" in data[:65536]:
            raise DocumentError("Binary content in a text file")
    return ext


def extract_text(ext: str, data: bytes) -> str:
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    text = data.decode("utf-8", errors="replace")
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return ""
        header = rows[0]
        return "\n".join("; ".join(f"{h}: {v}" for h, v in zip(header, r)) for r in rows[1:]) or ", ".join(header)
    return text


def chunk_text(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    text = "\n".join(line.rstrip() for line in text.splitlines())
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in paras:
        while len(p) > size:
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(p[:size])
            p = p[size - overlap:]
        if len(cur) + len(p) + 2 <= size:
            cur = f"{cur}\n\n{p}" if cur else p
        else:
            if cur:
                chunks.append(cur)
            tail = cur[-overlap:] if cur and overlap else ""
            cur = f"{tail}\n\n{p}" if tail else p
    if cur:
        chunks.append(cur)
    return chunks
